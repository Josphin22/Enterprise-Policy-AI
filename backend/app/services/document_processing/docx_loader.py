"""
DocxLoader extracts structured text, headings, sections, and tables from DOCX policy manuals.
Uses python-docx and formats tables into clean markdown pipe-delimited tables.
"""
import logging
from pathlib import Path
from typing import List, Optional
import docx
from app.services.document_processing.base_loader import BaseDocumentLoader, ExtractedBlock
from app.services.document_processing.table_extractor import table_extractor
from app.services.document_processing.section_detector import section_detector
from app.utils.errors import CorruptedDocumentError

logger = logging.getLogger("enterprise_rag.document_processing.docx")


class DocxLoader(BaseDocumentLoader):
    """
    Extracts structured text from DOCX policy manuals using python-docx.
    Captures paragraphs, headings, and table rows with section tracking.
    """

    def load(self, file_path: Path, **kwargs) -> List[ExtractedBlock]:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX document not found at: {file_path}")

        extracted_blocks: List[ExtractedBlock] = []

        try:
            try:
                doc = docx.Document(str(file_path))
            except Exception as open_err:
                logger.error(f"Failed to open DOCX file '{file_path.name}': {open_err}")
                raise CorruptedDocumentError(
                    f"Corrupted or unreadable DOCX document '{file_path.name}': {open_err}"
                ) from open_err

            current_section: Optional[str] = None
            accumulated_paragraphs: List[str] = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue

                style_name = paragraph.style.name.lower() if paragraph.style else ""
                is_heading_style = "heading" in style_name or "title" in style_name
                detected_heading = section_detector.detect_heading(text)

                if is_heading_style or detected_heading:
                    heading_title = detected_heading or text
                    # Flush previously accumulated paragraphs as a section block
                    if accumulated_paragraphs:
                        block_text = "\n\n".join(accumulated_paragraphs)
                        extracted_blocks.append(
                            ExtractedBlock(
                                text=block_text,
                                page_number=None,
                                section=current_section,
                                metadata={
                                    "filename": file_path.name,
                                    "file_type": "docx",
                                    "section": current_section,
                                },
                            )
                        )
                        accumulated_paragraphs = []

                    current_section = heading_title
                    accumulated_paragraphs.append(heading_title)
                else:
                    accumulated_paragraphs.append(text)

            # Flush any remaining paragraphs
            if accumulated_paragraphs:
                block_text = "\n\n".join(accumulated_paragraphs)
                extracted_blocks.append(
                    ExtractedBlock(
                        text=block_text,
                        page_number=None,
                        section=current_section,
                        metadata={
                            "filename": file_path.name,
                            "file_type": "docx",
                            "section": current_section,
                        },
                    )
                )

            # Extract tables into markdown representation
            for table_idx, table in enumerate(doc.tables, start=1):
                raw_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    if any(c for c in row_cells):
                        raw_rows.append(row_cells)

                if raw_rows:
                    md_table = table_extractor.format_markdown_table(raw_rows)
                    if md_table:
                        table_section = f"{current_section or 'General'} - Table {table_idx}"
                        extracted_blocks.append(
                            ExtractedBlock(
                                text=md_table,
                                page_number=None,
                                section=table_section,
                                metadata={
                                    "filename": file_path.name,
                                    "file_type": "docx",
                                    "is_table": True,
                                    "table_index": table_idx,
                                },
                            )
                        )

            logger.info(f"DocxLoader extracted {len(extracted_blocks)} structural blocks from {file_path.name}")

        except CorruptedDocumentError:
            raise
        except Exception as exc:
            logger.error(f"Failed to parse DOCX file '{file_path}': {exc}", exc_info=True)
            raise CorruptedDocumentError(f"Corrupted or unreadable DOCX document: {exc}") from exc

        return extracted_blocks


docx_loader = DocxLoader()
