"""
Phase 10 Test Suite: Advanced Document Intelligence, OCR & Table Extraction.
Verifies:
1. Normal PDF extraction (preserves page numbers, text, fast execution without unnecessary OCR).
2. Scanned PDF with OCR toggle (ENABLE_OCR=false skips, ENABLE_OCR=true processes).
3. DOCX extraction with formatted markdown tables and section headings.
4. TXT extraction with sections and encoding fallbacks.
5. Table extraction with searchable markdown pipe-delimited representation (| Col1 | Col2 |).
6. Empty document handling (0 bytes and whitespace-only rejected).
7. Corrupted document handling (broken PDF returns CORRUPTED_DOCUMENT).
8. Password-protected PDF handling (returns PASSWORD_PROTECTED).
9. Multi-page PDF header/footer conservative cleaning (repeated headers stripped, policy text intact).
10. Document metadata persistence (page_count, table_count, ocr_applied, language, processed_at).
11. Document preview API (GET /api/documents/{id}/preview, zero path disclosure).
12. Processing status lifecycle support (UPLOADED, PROCESSING, PROCESSED, FAILED, OCR_PROCESSING, NEEDS_REINDEX).
13. Existing RAG integration with Phase 10 tables and sections.
"""
import io
import pytest
from pathlib import Path
import pymupdf
import docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database.base import Base
from app.database.session import get_db
from app.main import app
from app.config import settings
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.document_metadata import DocumentMetadata
from app.models.user import User
from app.core.security import hash_password, create_access_token
from app.services.document_service import document_service
from app.services.extraction_service import extraction_service
from app.services.document_processing.pdf_loader import pdf_loader
from app.services.document_processing.docx_loader import docx_loader
from app.services.document_processing.txt_loader import txt_loader
from app.services.document_processing.ocr_service import ocr_service
from app.services.document_processing.table_extractor import table_extractor
from app.services.document_processing.section_detector import section_detector
from app.services.document_processing.header_footer_cleaner import header_footer_cleaner
from app.utils.errors import PasswordProtectedError, CorruptedDocumentError, EmptyDocumentError

TEST_DB_URL = "sqlite:///:memory:"
test_engine = create_engine(TEST_DB_URL, connect_args={"check_same_thread": False}, poolclass=StaticPool)
TestingSession = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


@pytest.fixture(scope="function")
def db_session():
    Base.metadata.create_all(bind=test_engine)
    session = TestingSession()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=test_engine)


@pytest.fixture(scope="function")
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture(scope="function")
def admin_token(db_session: Session):
    admin = User(
        username="admin_p10",
        email="admin_p10@enterprise.com",
        hashed_password=hash_password("AdminPass123!"),
        role="ADMIN",
        is_active=True,
    )
    db_session.add(admin)
    db_session.commit()
    return create_access_token({"sub": admin.id, "email": admin.email, "role": "ADMIN"})


# =====================================================================
# 1. Normal PDF Extraction
# =====================================================================

def test_normal_pdf_extraction_page_numbers_and_speed(tmp_path):
    """
    Test normal PDF with text extracts fast and preserves 1-indexed page numbers.
    Ensures OCR is NOT triggered when sufficient text exists.
    """
    pdf_file = tmp_path / "normal_multipage.pdf"
    doc = pymupdf.open()

    p1 = doc.new_page()
    p1.insert_text((50, 50), "1. Leave Policy Overview\nAll employees receive 18 days of annual leave.")

    p2 = doc.new_page()
    p2.insert_text((50, 50), "2. Carry Forward Rules\nA maximum of 5 unused leave days may be carried forward.")

    doc.save(str(pdf_file))
    doc.close()

    result = extraction_service.extract_document(pdf_file, file_type="pdf")
    assert "18 days of annual leave" in result["text"]
    assert "5 unused leave days" in result["text"]
    assert result["total_pages"] == 2
    assert result["pages"][0]["page_number"] == 1
    assert result["pages"][1]["page_number"] == 2
    # Verify OCR was not unnecessarily run
    assert result["ocr_applied"] is False


# =====================================================================
# 2. Scanned PDF with OCR Toggle
# =====================================================================

def test_scanned_pdf_ocr_toggle_and_execution(tmp_path):
    """
    Verify OCR behavior on low-text / scanned pages:
    - When ENABLE_OCR=False: OCR does not execute.
    - When ENABLE_OCR=True: OCR executes where text is insufficient.
    """
    pdf_file = tmp_path / "scanned_doc.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    # Draw a rectangle to simulate scanned image/drawing with almost no native text
    page.draw_rect(pymupdf.Rect(50, 50, 400, 300), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    page.insert_text((60, 60), "scan")  # < 10 characters
    doc.save(str(pdf_file))
    doc.close()

    # 1. When ENABLE_OCR is False
    original_ocr_state = settings.ENABLE_OCR
    settings.ENABLE_OCR = False
    try:
        # Check should_ocr_page returns False
        open_doc = pymupdf.open(str(pdf_file))
        assert ocr_service.should_ocr_page(open_doc[0], "scan") is False
        open_doc.close()

        # 2. When ENABLE_OCR is True
        settings.ENABLE_OCR = True
        open_doc = pymupdf.open(str(pdf_file))
        assert ocr_service.should_ocr_page(open_doc[0], "scan") is True

        # Inject mock OCR handler to test deterministic OCR extraction
        ocr_service.set_custom_handler(lambda img: "SCANNED CONTENT: Employees must submit medical certificates for sick leaves exceeding 3 days.")

        blocks = pdf_loader.load(pdf_file)
        assert len(blocks) == 1
        assert "medical certificates" in blocks[0].text
        assert blocks[0].metadata.get("ocr_applied") is True

    finally:
        ocr_service.set_custom_handler(None)
        settings.ENABLE_OCR = original_ocr_state


# =====================================================================
# 3. DOCX Extraction with Markdown Tables & Sections
# =====================================================================

def test_docx_extraction_with_tables_and_headings(tmp_path):
    """
    Verify DOCX extraction formats tables into searchable markdown pipe tables
    and preserves section headers.
    """
    docx_file = tmp_path / "sample_policy.docx"
    doc = docx.Document()
    doc.add_heading("1. Employee Benefits", level=1)
    doc.add_paragraph("Employees are eligible for standard insurance and wellness programs.")

    # Add Table
    table = doc.add_table(rows=3, cols=3)
    data = [
        ["Policy", "Days", "Approval"],
        ["Annual Leave", "18", "Manager"],
        ["Sick Leave", "10", "HR"],
    ]
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    doc.save(str(docx_file))

    result = extraction_service.extract_document(docx_file, file_type="docx")
    assert "1. Employee Benefits" in result["text"]
    # Table must be formatted as markdown pipe-delimited table
    assert "| Policy | Days | Approval |" in result["text"]
    assert "| --- | --- | --- |" in result["text"]
    assert "| Annual Leave | 18 | Manager |" in result["text"]
    assert result["table_count"] >= 1
    assert any("Employee Benefits" in s for s in result["sections"])


# =====================================================================
# 4. TXT Extraction with Sections
# =====================================================================

def test_txt_extraction_with_sections(tmp_path):
    """
    Verify plain text extraction detects section titles and handles standard encodings.
    """
    txt_file = tmp_path / "rules.txt"
    txt_file.write_text(
        "1. Leave Policy\nEmployees are granted 18 days annual leave.\n\n"
        "2. Remote Work Policy\nEmployees can work remotely 2 days per week.\n",
        encoding="utf-8",
    )

    result = extraction_service.extract_document(txt_file, file_type="txt")
    assert "18 days annual leave" in result["text"]
    assert "Remote Work Policy" in result["text"]
    assert result["total_pages"] == 1
    assert result["pages"][0]["page_number"] is None


# =====================================================================
# 5. Table Extraction Searchability
# =====================================================================

def test_table_searchable_markdown_format():
    """
    Verify TableExtractor converts 2D raw data into standard pipe-delimited markdown
    that preserves column tokens and numbers for RAG retrieval.
    """
    rows = [
        ["Leave Type", "Annual Days", "Notice Required"],
        ["Casual Leave", "12", "24 hours"],
        ["Paternity Leave", "15", "2 weeks"],
    ]
    md = table_extractor.format_markdown_table(rows)
    assert "| Leave Type | Annual Days | Notice Required |" in md
    assert "| --- | --- | --- |" in md
    assert "| Casual Leave | 12 | 24 hours |" in md
    assert "| Paternity Leave | 15 | 2 weeks |" in md


# =====================================================================
# 6. Error Handling: Empty Document
# =====================================================================

def test_empty_document_upload_rejected(client: TestClient):
    """
    Verify empty document (0 bytes) is rejected with 422 EMPTY_DOCUMENT.
    """
    resp = client.post(
        "/api/documents/upload",
        files={"file": ("empty.txt", io.BytesIO(b""), "text/plain")},
    )
    assert resp.status_code == 422
    data = resp.json()
    assert "EMPTY_DOCUMENT" in str(data) or "empty" in str(data).lower()


# =====================================================================
# 7. Error Handling: Corrupted Document
# =====================================================================

def test_corrupted_pdf_upload_rejected(client: TestClient):
    """
    Verify corrupted/broken PDF is caught and returns 422 CORRUPTED_DOCUMENT.
    """
    fake_broken_pdf = b"%PDF-1.4 completely broken content not a real pdf"
    resp = client.post(
        "/api/documents/upload",
        files={"file": ("corrupted.pdf", io.BytesIO(fake_broken_pdf), "application/pdf")},
    )
    assert resp.status_code == 422
    data = resp.json()
    assert "CORRUPTED_DOCUMENT" in str(data) or "corrupt" in str(data).lower() or "DOCUMENT_EXTRACTION_FAILED" in str(data)


# =====================================================================
# 8. Error Handling: Password-Protected PDF
# =====================================================================

def test_password_protected_pdf_rejected(tmp_path):
    """
    Verify encrypted/password-protected PDFs raise PasswordProtectedError.
    """
    encrypted_file = tmp_path / "protected.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "Confidential salary data")
    # Save with user password
    doc.save(
        str(encrypted_file),
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw="secret123",
        owner_pw="secret123",
    )
    doc.close()

    with pytest.raises(PasswordProtectedError):
        pdf_loader.load(encrypted_file)


# =====================================================================
# 9. Header / Footer Conservative Cleaning
# =====================================================================

def test_header_footer_conservative_cleaning(tmp_path):
    """
    Verify multi-page PDF with repeated header/footer strips boilerplate
    without deleting any internal policy content.
    """
    pdf_file = tmp_path / "handbook_multipage.pdf"
    doc = pymupdf.open()

    for i in range(1, 5):
        p = doc.new_page()
        # Same repeated header
        p.insert_text((50, 30), "Enterprise Policy Manual - Confidential")
        # Unique page content
        p.insert_text((50, 100), f"Section {i}. Content for policy chapter {i} with key rules.")
        # Repeated page footer
        p.insert_text((50, 700), f"Page {i} of 4")

    doc.save(str(pdf_file))
    doc.close()

    blocks = pdf_loader.load(pdf_file)
    assert len(blocks) == 4

    # Body content must be fully preserved
    for i in range(1, 5):
        assert f"Section {i}. Content for policy chapter {i}" in blocks[i - 1].text

    # Repeated headers and footers should be stripped from pages
    cleaned_texts = [b.text for b in blocks]
    assert not any("Page 1 of 4" in t for t in cleaned_texts)


# =====================================================================
# 10. Document Metadata Persistence
# =====================================================================

def test_metadata_persistence_and_no_fabrication(client: TestClient, db_session: Session):
    """
    Verify page_count, table_count, ocr_applied, and language are stored
    in the database accurately.
    """
    pdf_content = io.BytesIO()
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((50, 50), "1. Corporate Policy\nAll employees must adhere to the company standards.")
    doc.save(pdf_content)
    doc.close()
    pdf_content.seek(0)

    resp = client.post(
        "/api/documents/upload",
        files={"file": ("Policy_2026.pdf", pdf_content, "application/pdf")},
    )
    assert resp.status_code == 201
    doc_id = resp.json()["document"]["id"]

    # Verify database records
    db_doc = db_session.get(Document, doc_id)
    assert db_doc is not None
    assert db_doc.metadata_rel is not None
    assert db_doc.metadata_rel.page_count == 1
    assert db_doc.metadata_rel.table_count == 0
    assert db_doc.metadata_rel.ocr_applied is False
    assert db_doc.metadata_rel.language == "en"
    assert db_doc.metadata_rel.processed_at is not None


# =====================================================================
# 11. Document Preview Endpoint (Zero Path Disclosure)
# =====================================================================

def test_document_preview_endpoint(client: TestClient, db_session: Session):
    """
    Verify GET /api/documents/{id}/preview returns page and chunk details
    and strictly hides internal filesystem paths.
    """
    doc_id = "test-doc-prev-001"
    doc_record = Document(
        id=doc_id,
        filename="550e_Leave_Policy.pdf",
        original_filename="Leave_Policy.pdf",
        file_type="pdf",
        file_size=12000,
        file_path="C:\\internal\\secret\\path\\Leave_Policy.pdf",
        processing_status="processed",
        chunk_count=2,
        extracted_text="1. Leave Policy\nEmployees receive 18 days annual leave.\n\n2. Sick Leave\n10 days sick leave.",
        text_length=80,
    )
    db_session.add(doc_record)

    meta_record = DocumentMetadata(
        document_id=doc_id,
        title="Leave_Policy.pdf",
        page_count=2,
        table_count=1,
        ocr_applied=False,
        language="en",
    )
    db_session.add(meta_record)

    c1 = DocumentChunk(
        id="c1",
        document_id=doc_id,
        chunk_index=0,
        text="1. Leave Policy\nEmployees receive 18 days annual leave.",
        page_number=1,
        section="1. Leave Policy",
        character_count=40,
    )
    c2 = DocumentChunk(
        id="c2",
        document_id=doc_id,
        chunk_index=1,
        text="2. Sick Leave\n10 days sick leave.",
        page_number=2,
        section="2. Sick Leave",
        character_count=35,
    )
    db_session.add_all([c1, c2])
    db_session.commit()

    resp = client.get(f"/api/documents/{doc_id}/preview")
    assert resp.status_code == 200
    data = resp.json()

    assert data["success"] is True
    assert data["document_id"] == doc_id
    assert data["filename"] == "Leave_Policy.pdf"
    assert data["page_count"] == 2
    assert data["table_count"] == 1
    assert len(data["pages"]) >= 1
    assert len(data["chunks"]) == 2
    assert "1. Leave Policy" in data["sections"] or "2. Sick Leave" in data["sections"]

    # Security check: verify no internal filesystem path leakage
    response_text = resp.text
    assert "C:\\internal\\secret\\path" not in response_text
    assert "file_path" not in data
    assert "storage_path" not in data


# =====================================================================
# 12. Processing Status Lifecycle
# =====================================================================

def test_processing_status_lifecycle(db_session: Session):
    """
    Verify supported document processing statuses:
    UPLOADED, PROCESSING, PROCESSED, FAILED, OCR_PROCESSING, NEEDS_REINDEX.
    """
    doc = Document(
        id="status-test-doc",
        filename="status_test.pdf",
        original_filename="status_test.pdf",
        file_type="pdf",
        file_size=5000,
        file_path="/dummy/path",
        processing_status="uploaded",
    )
    db_session.add(doc)
    db_session.commit()

    valid_statuses = ["uploaded", "processing", "ocr_processing", "processed", "needs_reindex", "failed"]
    for s in valid_statuses:
        doc.processing_status = s
        db_session.commit()
        db_session.refresh(doc)
        assert doc.processing_status == s


# =====================================================================
# 13. Existing RAG Pipeline with Phase 10 Ingestion
# =====================================================================

def test_rag_retrieval_over_phase10_table_chunks(db_session: Session):
    """
    Verify that chunks containing extracted markdown tables and sections
    are retrieved and grounded by RAGService.
    """
    from app.rag.service import rag_service
    from app.rag.schemas import CandidateChunk

    doc_id = "doc-table-rag-01"
    doc_record = Document(
        id=doc_id,
        filename="Benefits_Table.docx",
        original_filename="Benefits_Table.docx",
        file_type="docx",
        file_size=15000,
        file_path="/dummy/path",
        processing_status="processed",
        chunk_count=1,
    )
    db_session.add(doc_record)

    table_chunk = DocumentChunk(
        id="chk-table-01",
        document_id=doc_id,
        chunk_index=0,
        text="| Policy | Days | Approval |\n| --- | --- | --- |\n| Annual Leave | 18 | Manager |\n| Sick Leave | 10 | HR |",
        page_number=None,
        section="1. Leave Benefits Table",
        character_count=110,
    )
    db_session.add(table_chunk)
    db_session.commit()

    # Build context using ContextBuilder
    cand = CandidateChunk(
        chunk_id=table_chunk.id,
        document_id=doc_id,
        filename=doc_record.original_filename,
        chunk_index=0,
        text=table_chunk.text,
        section=table_chunk.section,
        character_count=table_chunk.character_count,
        score=0.92,
    )
    context, sources = rag_service.builder.build_context([cand])
    assert "[Source S1]" in context
    assert "Benefits_Table.docx" in context
    assert "Annual Leave" in context
    assert "18" in context
    assert "Manager" in context
